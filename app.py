import datetime
import io
import os
import re
import zipfile
from docx import Document
from docx.oxml.ns import qn
from docx.shared import Inches, Pt
import pandas as pd
import qrcode
from PIL import Image
import requests
import streamlit as st
from streamlit_drawable_canvas import st_canvas

# ================= 1. 页面配置与初始化 =================
st.set_page_config(
    page_title="岗位职业危害告知书在线签收平台",
    layout="centered",
    initial_sidebar_state="expanded",
)

# 注入华文宋体全局样式与隐藏默认元素
hide_streamlit_style = """
    <style>
    #MainMenu {visibility: hidden;}
    footer {visibility: hidden;}
    header {visibility: hidden;}
    html, body, [class*="css"] {
        font-family: "华文宋体", SimSun, serif;
    }
    </style>
    """
st.markdown(hide_streamlit_style, unsafe_allow_html=True)

# ================= 2. 4位访问密码拦截验证 =================
if "authenticated" not in st.session_state:
  st.session_state.authenticated = False

if not st.session_state.authenticated:
  col_l, col_t = st.columns([1, 6])
  with col_l:
    try:
      st.image("logo.png", width=110)
    except Exception:
      st.image(
          "https://upload.wikimedia.org/wikipedia/commons/thumb/c/c5/Ikea_logo.svg/800px-Ikea_logo.svg.png",
          width=110,
      )
  with col_t:
    st.markdown("## 岗位职业危害告知书在线签收平台")

  st.markdown("---")
  st.info("🔒 本系统为内部合规平台，请输入 **4位访问密码** 进入系统。")

  with st.form("password_form"):
    pwd_input = st.text_input(
        "请输入 4 位访问密码：", type="password", max_chars=4
    )
    submit_pwd = st.form_submit_button("进入系统", use_container_width=True)

  if submit_pwd:
    correct_pwd = st.secrets.get("APP_PASSWORD", "8888")
    if pwd_input == correct_pwd:
      st.session_state.authenticated = True
      st.rerun()
    else:
      st.error("❌ 密码错误，请重新输入！")

  st.markdown("---")
  st.markdown(
      "<div style='text-align: center; color: gray; font-size: 14px;'>"
      "本系统为内部合规平台，严禁商业用途 | 开发者：陈野菲"
      "</div>",
      unsafe_allow_html=True,
  )
  st.stop()


# ================= 3. 百度网盘自动上传函数（带 OAuth2 自动刷新） =================
def refresh_baidu_access_token():
  try:
    client_id = st.secrets.get("BAIDU_CLIENT_ID", "")
    client_secret = st.secrets.get("BAIDU_CLIENT_SECRET", "")
    refresh_token = st.secrets.get("BAIDU_REFRESH_TOKEN", "")
    if not client_id or not client_secret or not refresh_token:
      return None
    token_url = "https://pan.baidu.com/oauth/2.0/token"
    params = {
        "grant_type": "refresh_token",
        "refresh_token": refresh_token,
        "client_id": client_id,
        "client_secret": client_secret,
    }
    response = requests.get(token_url, params=params)
    res_data = response.json()
    return res_data.get("access_token")
  except Exception:
    return None


def upload_to_baidu_netdisk_with_auto_refresh(file_bytes, remote_filename):
  access_token = st.secrets.get("BAIDU_ACCESS_TOKEN", "")
  if not access_token:
    return (
        False,
        "未配置网盘凭证，文件已在本地生成并可通过网页下载/ZIP打包保存。",
    )

  sub_folder = "EHS签字档案"
  target_path = f"/apps/慧瑞EHS合规档案/{sub_folder}/{remote_filename}"

  def send_upload_request(token):
    upload_url = f"https://pan.baidu.com/rest/2.0/xpan/file?method=upload&access_token={token}&path={target_path}&uploadid=&file=1"
    files = {"file": (remote_filename, file_bytes)}
    return requests.post(upload_url, files=files).json()

  result = send_upload_request(access_token)
  if "errno" in result and result["errno"] in [110, 111]:
    new_token = refresh_baidu_access_token()
    if new_token:
      result = send_upload_request(new_token)
    else:
      return False, "Token 已过期且自动刷新失败。"

  if "errno" in result and result["errno"] == 0:
    return True, f"成功同步至网盘：/apps/慧瑞EHS合规档案/{sub_folder}/"
  else:
    return False, f"网盘上传失败: {result.get('error_msg', '未知错误')}"


# ================= 4. 侧边栏：Logo与微信分享 =================
with st.sidebar:
  try:
    st.image("logo.png", width=160)
  except Exception:
    st.image(
        "https://upload.wikimedia.org/wikipedia/commons/thumb/c/c5/Ikea_logo.svg/800px-Ikea_logo.svg.png",
        width=160,
    )

  st.markdown("### 📱 微信扫码与分享")
  st.write("已自动关联您的云端网址，二维码将实时更新供手机扫码填报。")

  app_url = st.text_input(
      "应用公网链接 (URL)", value="https://occupational-check-sign.streamlit.app"
  )

  if app_url:
    qr = qrcode.make(app_url)
    img_buffer = io.BytesIO()
    qr.save(img_buffer, format="PNG")
    st.image(
        Image.open(img_buffer), caption="微信扫码快速查阅与签收", width=160
    )
    st.info(
        "💡 **提示**：将上方链接复制并发送至微信工作群，员工即可手机端完成告知书签收。"
    )

# ================= 5. 主界面逻辑（Logo在左侧，主标题单独一行） =================
col_logo, col_title = st.columns([1, 6])
with col_logo:
  try:
    st.image("logo.png", width=110)
  except Exception:
    st.image(
        "https://upload.wikimedia.org/wikipedia/commons/thumb/c/c5/Ikea_logo.svg/800px-Ikea_logo.svg.png",
        width=110,
    )
with col_title:
  st.markdown("## 岗位职业危害告知书在线签收平台")

with st.expander("📢 依据《职业病防治法》及国家规范的法定要求说明", expanded=False):
  st.info("""
    本系统及告知流程严格依据以下国家法律法规和标准建设执行：
    * **《中华人民共和国职业病防治法》及《用人单位职业病危害告知与警示标识管理规范》（安总厅安健〔2014〕111号）**：
      * **第 六 条**：产生职业病危害的用人单位应将工作过程中可能接触的危害因素种类、程度、后果、防护设施、个人防护用品、健康检查和待遇等**如实告知劳动者，不得隐瞒或者欺骗**。
      * **第 七 条**：劳动合同时应当写明职业病危害相关内容；格式合同文本内容不完善的，**应以合同附件形式签署职业病危害告知书**。
      * **第 九 条**：用人单位应对劳动者进行**上岗前及定期职业卫生培训**，经考试合格后方可上岗作业。
      * **第 十 一 条**：应按照规定组织接触职业病危害作业的劳动者进行**上岗前、在岗期间和离岗时的职业健康检查**，并将结果书面告知劳动者。
    """)

st.markdown("---")

# 基础信息录入
st.subheader("1. 员工身份核验")
col1, col2 = st.columns(2)
with col1:
  emp_name = st.text_input("员工姓名 (必填)：")
with col2:
  emp_id = st.text_input(
      "身份证号 (必填，18位)：",
      help="请输入标准的 18 位中国居民身份证号码",
  )

st.write("---")
st.markdown("### 📂 选择并查阅岗位职业危害告知书")

departments = ["技术部", "品管部", "生产部"]
stages = ["上岗前", "在岗期间"]

col_d, col_s = st.columns(2)
with col_d:
  selected_dept = st.selectbox(
      "👉 请选择所属部门：", options=["-- 请选择部门 --"] + departments, index=0
  )
with col_s:
  selected_stage = st.selectbox(
      "👉 请选择告知阶段：", options=["-- 请选择阶段 --"] + stages, index=0
  )


# 智能模糊匹配 attachments 文件夹中的 .docx 文件
def find_attachment_file(folder, dept, stage):
  if not os.path.exists(folder):
    return None, None
  for filename in os.listdir(folder):
    if dept in filename and stage in filename and filename.lower().endswith(".docx"):
      return os.path.join(folder, filename), filename
  return None, None


ATTACHMENT_DIR = "attachments"
file_path, matched_filename = None, None

if selected_dept != "-- 请选择部门 --" and selected_stage != "-- 请选择阶段 --":
  file_path, matched_filename = find_attachment_file(
      ATTACHMENT_DIR, selected_dept, selected_stage
  )

if file_path and os.path.exists(file_path):
  st.success(f"✅ 成功找到专属告知书文件：【{matched_filename}】")
  with open(file_path, "rb") as fr:
    hazard_bytes = fr.read()

  st.download_button(
      label=f"📥 点击下载并查阅《{matched_filename}》",
      data=hazard_bytes,
      file_name=matched_filename,
      mime=(
          "application/vnd.openxmlformats-officedocument.wordprocessingml.document"
      ),
      use_container_width=True,
  )
else:
  if selected_dept != "-- 请选择部门 --" and selected_stage != "-- 请选择阶段 --":
    st.warning(
        f"⚠️ 未在 'attachments' 文件夹中检测到匹配的 `.docx` 文件（当前检索条件：【{selected_dept}】+【{selected_stage}】）。请确认已将 Word 模板转为 `.docx` 格式并上传。"
    )
  else:
    st.info("💡 请先在上方完整选择您的部门和告知阶段以加载告知书。")

c_hazard = st.checkbox(
    "【须确认】本人已阅读并充分了解上述岗位职业危害告知书的内容，已知悉各项职业危害因素及防护要求，承诺在工作中严格落实。"
)

# ================= 6. 手写签名与手写日期栏（并排双画布） =================
current_date_str = datetime.date.today().strftime("%Y年%m月%d日")

st.write("---")
st.subheader("✍️ 2. 员工手写签名与手写日期栏")
st.markdown(
    f"**请在左侧手写签名，并在右侧手写日期（注：当前系统日期为"
    f" {current_date_str}，请按此手写日期）：**"
)

col_sig, col_date = st.columns(2)
with col_sig:
  st.markdown("**手写签名：**")
  canvas_result = st_canvas(
      stroke_width=4,
      stroke_color="#000000",
      background_color="#F8F9FA",
      height=200,
      width=320,
      drawing_mode="freedraw",
      key="canvas_sig",
      return_image_data=True,
  )
with col_date:
  st.markdown("**手写日期栏（请手写当前日期）：**")
  canvas_date_result = st_canvas(
      stroke_width=3,
      stroke_color="#000000",
      background_color="#F8F9FA",
      height=200,
      width=320,
      drawing_mode="freedraw",
      key="canvas_date",
      return_image_data=True,
  )


# ================= 7. 辅助函数：向 Word 模板文末追加签名与日期 =================
def append_signature_to_docx(
    template_path, default_title, sig_image_io, date_image_io
):
  if template_path and os.path.exists(template_path):
    try:
      doc = Document(template_path)
    except Exception:
      doc = Document()
      doc.add_heading(default_title, level=1)
      doc.add_paragraph("（提示：模板文件读取异常，此为生成的标准确认单）")
  else:
    doc = Document()
    doc.add_heading(default_title, level=1)
    doc.add_paragraph("（提示：未找到对应的 .docx 模板文件）")

  for p in doc.paragraphs:
    for r in p.runs:
      r.font.name = "华文宋体"
      r.font.element.rPr.rFonts.set(qn("w:eastAsia"), "华文宋体")

  p_line = doc.add_paragraph("--------------------------------------------------")
  p_line.paragraph_format.space_before = Pt(5)
  p_line.paragraph_format.space_after = Pt(5)

  p_confirm = doc.add_paragraph()
  run_c = p_confirm.add_run(
      f"【员工签收确认】 姓名：{emp_name} | 身份证号：{emp_id}\n"
      f"签收时间：{datetime.datetime.now().strftime('%Y-%m-%d %H:%M:%S')}\n"
      f"本人已仔细阅读并充分理解上述岗位职业危害告知书内容，承诺严格遵守各项安全防护规定。"
  )
  run_c.font.name = "华文宋体"
  run_c.font.size = Pt(10.5)
  run_c.font.element.rPr.rFonts.set(qn("w:eastAsia"), "华文宋体")

  table = doc.add_table(rows=1, cols=2)
  table.autofit = False

  cell_sig = table.cell(0, 0)
  p1 = cell_sig.paragraphs[0]
  r1 = p1.add_run("员工手写亲笔签名：\n")
  r1.font.name = "华文宋体"
  r1.font.size = Pt(10)
  r1.font.element.rPr.rFonts.set(qn("w:eastAsia"), "华文宋体")
  p1.add_run().add_picture(sig_image_io, width=Inches(1.8))
  sig_image_io.seek(0)

  cell_date = table.cell(0, 1)
  p2 = cell_date.paragraphs[0]
  r2 = p2.add_run("手写签署日期：\n")
  r2.font.name = "华文宋体"
  r2.font.size = Pt(10)
  r2.font.element.rPr.rFonts.set(qn("w:eastAsia"), "华文宋体")
  p2.add_run().add_picture(date_image_io, width=Inches(1.8))
  date_image_io.seek(0)

  buffer = io.BytesIO()
  doc.save(buffer)
  buffer.seek(0)
  return buffer


# ================= 8. 提交校验与生成档案 =================
if st.button(
    "📁 确认无误，一键签收告知书并生成合规档案", use_container_width=True
):
  is_canvas_empty = canvas_result.image_data is None or (
      canvas_result.json_data is not None
      and len(canvas_result.json_data.get("objects", [])) == 0
  )
  is_date_empty = canvas_date_result.image_data is None or (
      canvas_date_result.json_data is not None
      and len(canvas_date_result.json_data.get("objects", [])) == 0
  )

  id_pattern = re.compile(r"^\d{17}[\dXx]$")

  if not emp_name.strip() or not emp_id.strip():
    st.error("❌ 拦截 : 请完整填写【员工姓名】与【身份证号】！")
  elif not id_pattern.match(emp_id.strip()):
    st.error(
        "❌ 拦截 : 身份证号必须为严格的 **18 位**数字（末尾可为大写 X）！"
    )
  elif not file_path:
    st.error("❌ 拦截 : 请先选择正确的部门和阶段以加载对应的告知书文件！")
  elif not c_hazard:
    st.error("❌ 拦截 : 您必须勾选确认已阅读并知悉岗位职业危害告知书！")
  elif is_canvas_empty:
    st.warning("⚠️ 拦截 : 请在左侧画板完成手写签名后再提交。")
  elif is_date_empty:
    st.warning("⚠️ 拦截 : 请在右侧手写日期栏内完成手写日期后再提交！")
  else:
    st.success(
        "✅ 告知书签收成功！系统已成功生成您的专属带签名 Word 合规确认凭证。"
    )

    signature_img = Image.fromarray(
        canvas_result.image_data.astype("uint8"), "RGBA"
    )
    sig_io = io.BytesIO()
    signature_img.save(sig_io, format="PNG")
    sig_io.seek(0)

    date_img = Image.fromarray(
        canvas_date_result.image_data.astype("uint8"), "RGBA"
    )
    date_io = io.BytesIO()
    date_img.save(date_io, format="PNG")
    date_io.seek(0)

    # 生成带签名的 Word 告知书
    signed_hazard_buffer = append_signature_to_docx(
        file_path, f"职业危害告知书 - {selected_dept}（{selected_stage}）", sig_io, date_io
    )

    # 动态打包 ZIP
    zip_buffer = io.BytesIO()
    with zipfile.ZipFile(zip_buffer, "w", zipfile.ZIP_DEFLATED) as zip_file:
      # 1. 放入带签名的 Word 告知书
      hazard_filename_cloud = f"职业危害告知书_{selected_dept}_{selected_stage}_{emp_name}_{emp_id[-4:]}_已签字.docx"
      zip_file.writestr(
          hazard_filename_cloud, signed_hazard_buffer.getvalue()
      )

      # 2. 自动同步到百度网盘
      upload_to_baidu_netdisk_with_auto_refresh(
          signed_hazard_buffer.getvalue(), hazard_filename_cloud
      )

      # 3. 保存签名及日期原图
      img_byte_arr = io.BytesIO()
      signature_img.save(img_byte_arr, format="PNG")
      zip_file.writestr(
          f"手写签名原图_{emp_name}.png", img_byte_arr.getvalue()
      )

      date_byte_arr = io.BytesIO()
      date_img.save(date_byte_arr, format="PNG")
      zip_file.writestr(f"手写日期原图_{emp_name}.png", date_byte_arr.getvalue())

    zip_buffer.seek(0)

    st.markdown("---")
    st.success(
        "🎉 您的告知书签收档案已打包完毕，点击下方按钮即可下载保存！"
    )

    col_d1, col_d2 = st.columns(2)
    with col_d1:
      st.download_button(
          label="📄 下载带签名的告知书 (.docx)",
          data=signed_hazard_buffer.getvalue(),
          file_name=f"职业危害告知书_{selected_dept}_{selected_stage}_{emp_name}.docx",
          mime=(
              "application/vnd.openxmlformats-officedocument.wordprocessingml.document"
          ),
          use_container_width=True,
      )
    with col_d2:
      st.download_button(
          label="📥 一键打包下载全部档案 (.ZIP)",
          data=zip_buffer,
          file_name=f"危害告知书签收档案_{emp_name}.zip",
          mime="application/zip",
          use_container_width=True,
      )

    st.snow()

# ================= 9. 底部版权与开发者声明 =================
st.markdown("---")
st.markdown(
    "<div style='text-align: center; color: gray; font-size: 14px;'>"
    "本系统为内部合规平台，严禁商业用途 | 开发者：陈野菲"
    "</div>",
    unsafe_allow_html=True,
)
